#!/usr/bin/env python3
"""Converts the fitness Excel workbook into the SQLite seed database for the iOS app."""
import argparse
import sqlite3
import sys
import warnings
from collections import Counter
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path

import openpyxl

warnings.filterwarnings("ignore", message="Unknown extension is not supported")

PROJECT_ROOT = Path(__file__).resolve().parents[2]
DEFAULT_XLSX = PROJECT_ROOT / "current_artifacts" / "Data File - Single Line.xlsx"
DEFAULT_OUT = PROJECT_ROOT / "output" / "fitness.sqlite"
SCHEMA_PATH = PROJECT_ROOT / "src" / "shared" / "database" / "schema.sql"

MASTER_SHEET = "Master Sheet"
PLAN_SHEET = "V5Fullbody_JMK"
LIBRARY_SHEET = "Exercise to muslce & instructio"

TABLES = [
    "exercises", "plan_days", "plan_entries", "set_logs", "body_weights",
    "body_measurements", "measurement_targets", "calorie_entries", "settings",
]


class SeedError(Exception):
    pass


class MissingLibraryExercise(SeedError):
    pass


@dataclass
class Skip:
    sheet: str
    row: int
    table: str
    reason: str


@dataclass
class Report:
    counts: dict = field(default_factory=dict)
    skipped: list = field(default_factory=list)
    log_only_created: int = 0
    null_muscle_exercises: int = 0

    def summary(self):
        lines = ["Rows per table:"]
        lines += [f"  {t}: {self.counts[t]}" for t in TABLES]
        lines.append(f"Log-only exercises created: {self.log_only_created}")
        lines.append(f"Exercises with NULL muscle: {self.null_muscle_exercises}")
        lines.append(f"Skipped rows: {len(self.skipped)}")
        by_reason = Counter((s.table, s.reason) for s in self.skipped)
        lines += [f"  {n} x {table}: {reason}" for (table, reason), n in sorted(by_reason.items())]
        lines += [f"    {s.sheet} row {s.row}: {s.reason}" for s in self.skipped]
        return "\n".join(lines)


def clean_text(value):
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def to_iso_date(value):
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    text = clean_text(value)
    if text is None:
        return None
    try:
        return datetime.fromisoformat(text).date().isoformat()
    except ValueError:
        return None


def to_number(value):
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return float(value)
    text = clean_text(value)
    if text is None:
        return None
    try:
        return float(text)
    except ValueError:
        return None


def to_int(value):
    number = to_number(value)
    if number is None or not number.is_integer():
        return None
    return int(number)


def read_sheet(workbook, sheet_name, required_headers):
    """Yields (sheet_row_number, {header: value}) for every non-blank data row."""
    if sheet_name not in workbook.sheetnames:
        raise SeedError(f"Sheet not found: {sheet_name!r}")
    rows = workbook[sheet_name].iter_rows(values_only=True)
    headers = [clean_text(h) for h in next(rows, ())]
    missing = [h for h in required_headers if h not in headers]
    if missing:
        raise SeedError(f"Sheet {sheet_name!r} is missing columns: {missing}")
    wanted = {h: headers.index(h) for h in required_headers}
    for number, row in enumerate(rows, start=2):
        record = {h: (row[i] if i < len(row) else None) for h, i in wanted.items()}
        if any(v is not None and str(v).strip() != "" for v in record.values()):
            yield number, record


def load_library(workbook, report):
    library = {}
    for number, row in read_sheet(workbook, LIBRARY_SHEET, ["Exercise Name", "Muscle", "Instructions"]):
        name = clean_text(row["Exercise Name"])
        if name is None:
            report.skipped.append(Skip(LIBRARY_SHEET, number, "exercises", "missing exercise name"))
        elif name.lower() in library:
            # the schema forbids case-insensitive duplicates: the first row wins, later rows only fill its blanks
            first_name, muscle, instructions = library[name.lower()]
            merged = (first_name, muscle or clean_text(row["Muscle"]), instructions or clean_text(row["Instructions"]))
            library[name.lower()] = merged
            report.skipped.append(Skip(
                LIBRARY_SHEET, number, "exercises",
                f"duplicate library name {name!r} merged into {first_name!r} (first row wins, blanks filled)",
            ))
        else:
            library[name.lower()] = (name, clean_text(row["Muscle"]), clean_text(row["Instructions"]))
    return library


@dataclass
class MasterRows:
    exercise_sets: list = field(default_factory=list)  # (row, key, date, weight, reps, muscle)
    exercise_names: dict = field(default_factory=dict)  # key -> first-seen spelling
    muscle_votes: dict = field(default_factory=dict)  # key -> [(date, row, muscle)]
    body_weights: list = field(default_factory=list)
    body_measurements: list = field(default_factory=list)
    targets: dict = field(default_factory=dict)  # site key -> (date, row, site, inches)


def classify_master(workbook, report):
    master = MasterRows()
    headers = [
        "Category", "Date", "Muscle", "Measurement (in inches)",
        "Name of the exercise", "Repetitions performed", "Weight (In Kg)",
    ]

    def skip(number, table, reason):
        report.skipped.append(Skip(MASTER_SHEET, number, table, reason))

    for number, row in read_sheet(workbook, MASTER_SHEET, headers):
        category = clean_text(row["Category"])
        iso_date = to_iso_date(row["Date"])
        muscle = clean_text(row["Muscle"])
        if category == "Exercise":
            name = clean_text(row["Name of the exercise"])
            if name is not None:
                key = name.lower()
                master.exercise_names.setdefault(key, name)
                master.muscle_votes.setdefault(key, []).append((iso_date or "", number, muscle))
            weight = to_number(row["Weight (In Kg)"])
            reps = to_int(row["Repetitions performed"])
            problem = (
                "missing exercise name" if name is None
                else "missing or invalid date" if iso_date is None
                else "missing or invalid weight" if weight is None or weight < 0
                else "missing or invalid reps" if reps is None or reps < 0
                else None
            )
            if problem:
                skip(number, "set_logs", problem)
            else:
                master.exercise_sets.append((number, name.lower(), iso_date, weight, reps, muscle))
        elif category == "Weight":
            kg = to_number(row["Weight (In Kg)"])
            if iso_date is None:
                skip(number, "body_weights", "missing or invalid date")
            elif kg is None or kg <= 0:
                skip(number, "body_weights", "missing or invalid weight")
            else:
                master.body_weights.append((iso_date, kg))
        elif category in ("Body Measurement", "Target Measurement"):
            table = "body_measurements" if category == "Body Measurement" else "measurement_targets"
            inches = to_number(row["Measurement (in inches)"])
            if muscle is None:
                skip(number, table, "missing site")
            elif iso_date is None:
                skip(number, table, "missing or invalid date")
            elif inches is None or inches <= 0:
                skip(number, table, "missing or invalid inches")
            elif table == "body_measurements":
                master.body_measurements.append((iso_date, muscle, inches))
            else:
                current = master.targets.get(muscle.lower())
                if current is None or (iso_date, number) >= (current[0], current[1]):
                    master.targets[muscle.lower()] = (iso_date, number, muscle, inches)
        else:
            skip(number, "-", f"unknown category {category!r}")
    return master


def majority_muscle(votes):
    labelled = [(d, n, m) for d, n, m in votes if m is not None]
    if not labelled:
        return None
    counts = Counter(m for _, _, m in labelled)
    latest_for = {}
    for d, n, m in labelled:
        latest_for[m] = max(latest_for.get(m, (d, n)), (d, n))
    return max(counts, key=lambda m: (counts[m], latest_for[m]))


def build(xlsx_path, out_path, schema_path=SCHEMA_PATH):
    xlsx_path, out_path = Path(xlsx_path), Path(out_path)
    report = Report()
    workbook = openpyxl.load_workbook(xlsx_path, data_only=True, read_only=True)
    try:
        library = load_library(workbook, report)
        plan_rows = list(read_sheet(
            workbook, PLAN_SHEET,
            ["Exercise Number", "Day", "Name of the exercise", "Rep Type", "Sets to Perform", "Minreps", "Maxreps"],
        ))
        master = classify_master(workbook, report)
    finally:
        workbook.close()

    for number, row in plan_rows:
        name = clean_text(row["Name of the exercise"])
        if name is None or name.lower() not in library:
            raise MissingLibraryExercise(
                f"{PLAN_SHEET} row {number}: exercise {name!r} is not in the library sheet"
            )

    out_path.parent.mkdir(parents=True, exist_ok=True)
    if out_path.exists():
        out_path.unlink()

    conn = sqlite3.connect(out_path, isolation_level=None)
    try:
        conn.executescript(Path(schema_path).read_text())
        conn.execute("PRAGMA foreign_keys=ON")
        conn.execute("BEGIN")
        try:
            insert_all(conn, report, library, plan_rows, master)
            check_clean(conn)
            conn.execute("COMMIT")
        except BaseException:
            conn.execute("ROLLBACK")
            raise
        report.counts = {t: conn.execute(f"SELECT COUNT(*) FROM {t}").fetchone()[0] for t in TABLES}
        report.null_muscle_exercises = conn.execute(
            "SELECT COUNT(*) FROM exercises WHERE muscle IS NULL"
        ).fetchone()[0]
        conn.execute("VACUUM")
    finally:
        conn.close()
    return report


def insert_all(conn, report, library, plan_rows, master):
    ids = {}

    def add_exercise(name, muscle, instructions, source):
        cursor = conn.execute(
            "INSERT INTO exercises (name, muscle, instructions, source) VALUES (?, ?, ?, ?)",
            (name, muscle, instructions, source),
        )
        ids[name.lower()] = cursor.lastrowid

    for name, muscle, instructions in library.values():
        add_exercise(name, muscle, instructions, "library")

    for key, name in master.exercise_names.items():
        if key not in ids:
            add_exercise(name, majority_muscle(master.muscle_votes[key]), None, "log")
            report.log_only_created += 1

    entries = [plan_entry(number, row, ids) for number, row in plan_rows]
    day_ids = {}
    for day in sorted({entry[0] for entry in entries}):
        day_ids[day] = conn.execute("INSERT INTO plan_days (position) VALUES (?)", (day,)).lastrowid
    conn.executemany(
        """INSERT INTO plan_entries (day_id, position, exercise_id, sets, min_reps, max_reps, rep_type)
           VALUES (?, ?, ?, ?, ?, ?, ?)""",
        [(day_ids[day], *rest) for day, *rest in entries],
    )
    conn.executemany(
        "INSERT INTO set_logs (performed_on, exercise_id, weight_kg, reps, logged_muscle) VALUES (?, ?, ?, ?, ?)",
        [(d, ids[key], weight, reps, muscle) for _, key, d, weight, reps, muscle in master.exercise_sets],
    )
    conn.executemany("INSERT INTO body_weights (measured_on, weight_kg) VALUES (?, ?)", master.body_weights)
    conn.executemany(
        "INSERT INTO body_measurements (measured_on, site, inches) VALUES (?, ?, ?)", master.body_measurements
    )
    conn.executemany(
        "INSERT INTO measurement_targets (site, inches, set_on) VALUES (?, ?, ?)",
        [(site, inches, d) for d, _, site, inches in master.targets.values()],
    )


def plan_entry(number, row, ids):
    day, position = to_int(row["Day"]), to_int(row["Exercise Number"])
    sets, min_reps, max_reps = to_int(row["Sets to Perform"]), to_int(row["Minreps"]), to_int(row["Maxreps"])
    if None in (day, position, sets, min_reps, max_reps):
        raise SeedError(f"{PLAN_SHEET} row {number}: missing or non-integer day/position/sets/reps")
    exercise_id = ids[clean_text(row["Name of the exercise"]).lower()]
    return (day, position, exercise_id, sets, min_reps, max_reps, clean_text(row["Rep Type"]))


def check_clean(conn):
    violations = conn.execute("PRAGMA foreign_key_check").fetchall()
    if violations:
        raise SeedError(f"foreign_key_check failed: {violations[:5]}")
    result = conn.execute("PRAGMA integrity_check").fetchall()
    if result != [("ok",)]:
        raise SeedError(f"integrity_check failed: {result[:5]}")


def main(argv=None):
    parser = argparse.ArgumentParser(description="Build the fitness SQLite seed database from the workbook.")
    parser.add_argument("--xlsx", default=str(DEFAULT_XLSX), help="source workbook")
    parser.add_argument("--out", default=str(DEFAULT_OUT), help="output SQLite file (overwritten)")
    args = parser.parse_args(argv)
    try:
        report = build(args.xlsx, args.out)
    except SeedError as error:
        print(f"ERROR: {error}", file=sys.stderr)
        return 1
    print(f"Wrote {args.out}")
    print(report.summary())
    return 0


if __name__ == "__main__":
    sys.exit(main())
