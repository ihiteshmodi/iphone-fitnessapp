from datetime import datetime

import openpyxl

MASTER_HEADER = [
    "Category", "Date", "Muscle", "Measurement (in inches)",
    "Name of the exercise", "Repetitions performed", "Weight (In Kg)",
]
PLAN_HEADER = [
    "Week", "Exercise Number", "Day", "Name of the exercise", "Rep Type",
    "Sets to Perform", "Minreps", "Maxreps", "Instructions", "Muscle",
]
LIBRARY_HEADER = ["Exercise Name", "Muscle", "Instructions"]

D = datetime


def exercise_row(on, name, reps, weight, muscle=None):
    return ["Exercise", on, muscle, None, name, reps, weight]


def weight_row(on, kg):
    return ["Weight", on, None, None, None, None, kg]


def measurement_row(on, site, inches):
    return ["Body Measurement", on, site, inches, None, None, None]


def target_row(on, site, inches):
    return ["Target Measurement", on, site, inches, None, None, None]


def plan_row(day, position, name, sets=3, min_reps=8, max_reps=12, rep_type=None):
    return [1, position, day, name, rep_type, sets, min_reps, max_reps, "ignored", "ignored"]


class WorkbookSpec:
    """Collects rows for the three relevant sheets and writes a synthetic workbook."""

    def __init__(self):
        self.master = []
        self.plan = []
        self.library = []

    def write(self, path):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Master Sheet"
        for row in [MASTER_HEADER, *self.master]:
            ws.append(row)
        ws = wb.create_sheet("V5Fullbody_JMK")
        for row in [PLAN_HEADER, *self.plan]:
            ws.append(row)
        ws = wb.create_sheet("Exercise to muslce & instructio")
        for row in [LIBRARY_HEADER, *self.library]:
            ws.append(row)
        wb.save(path)
        return path
